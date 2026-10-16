"""
Final price lists — what each supplier will be paid, signed by SPAR.

At the approval cutoff (Admin > Settings, default 2:00 PM) on the price
submission day every price for that delivery date is settled: the supplier
can't answer SPAR's adjustments any more and Admin can't send new ones
(price_approval_service.expires_at). From then on the supplier sees their
final list for that date, item by item:

- an adjusted price they approved -> SPAR's adjusted price
- one they rejected with their own final price -> that price
- everything else (rejected without a price, never answered, never
  adjusted) -> the price they submitted

That's the same rule Supplier Orders and purchase orders use
(supplier_order_service._prices_from_rows) once the cutoff has passed.

The list carries SPAR's authorised signature — drawn or uploaded once by
Admin on Settings and stamped on every list automatically — and the
supplier's own e-signature from any price sheet they approved for the
date. Suppliers download it as Excel or PDF.
"""
import base64
import hashlib
import io
import json
from datetime import date, datetime, timedelta
from pathlib import Path

from sqlalchemy.orm import Session

from app.core.audit import write_audit_log
from app.core.errors import NotFoundError, PermissionDeniedError, ValidationFailedError
from app.models.price_revision import SupplierPriceRevision
from app.models.pricing import SupplierPrice
from app.models.product import Product, ProductCategory
from app.models.supplier import Supplier
from app.models.system import SystemSetting
from app.models.user import User
from app.services import price_approval_service as approvals

# SPAR's authorised signature. Kept in system_settings but outside
# settings_service's public keys — GET /settings is readable without login.
SIGNATURE_IMAGE = "final_price_signature_image"
SIGNATORY_NAME = "final_price_signatory_name"
SIGNATORY_TITLE = "final_price_signatory_title"

_SPAR_LOGO_PATH = Path(__file__).resolve().parent.parent / "assets" / "spar-logo.png"
_SPAR_GREEN = "157946"
# How many past delivery dates the supplier can pick from.
MAX_DATES = 60

# Why an item has the final price it has.
BASIS_APPROVED = "SPAR price approved"
BASIS_COUNTER = "Your final price"
BASIS_REJECTED = "Your price (SPAR price rejected)"
BASIS_EXPIRED = "Your price (not answered)"
BASIS_OWN = "Your price"


# ---- SPAR's signature ----


def get_signature(db: Session) -> dict:
    rows = {
        r.key: r
        for r in db.query(SystemSetting).filter(
            SystemSetting.key.in_((SIGNATURE_IMAGE, SIGNATORY_NAME, SIGNATORY_TITLE))
        )
    }
    image = rows.get(SIGNATURE_IMAGE)
    updated_by = db.get(User, image.updated_by) if image and image.updated_by else None
    return {
        "signature_image": image.value if image else None,
        "signatory_name": rows[SIGNATORY_NAME].value if SIGNATORY_NAME in rows else None,
        "signatory_title": rows[SIGNATORY_TITLE].value if SIGNATORY_TITLE in rows else None,
        "updated_at": image.updated_at if image else None,
        "updated_by_name": updated_by.username if updated_by else None,
    }


def set_signature(db: Session, admin: User, *, name: str, title: str | None, signature_image: str) -> dict:
    name = name.strip()
    if not name:
        raise ValidationFailedError("Enter the name of the person signing for SPAR.")
    approvals._check_signature_image(signature_image)
    for key, value in ((SIGNATURE_IMAGE, signature_image), (SIGNATORY_NAME, name), (SIGNATORY_TITLE, (title or "").strip())):
        row = db.get(SystemSetting, key)
        if row:
            row.value = value
            row.updated_by = admin.id
        else:
            db.add(SystemSetting(key=key, value=value, updated_by=admin.id))
    db.commit()
    write_audit_log(
        db,
        user_id=admin.id,
        role="ADMIN",
        action="FINAL_PRICE_SIGNATURE_SET",
        entity_type="system_setting",
        entity_id=None,
        description=f"Set SPAR's signature on final price lists (signatory: {name}).",
    )
    return get_signature(db)


# ---- The list ----


def _supplier_id(user: User) -> int:
    if not user.supplier_id:
        raise PermissionDeniedError("Only supplier accounts have final price lists.")
    return user.supplier_id


def final_dates(db: Session, supplier_user: User) -> dict:
    """Delivery dates whose prices are final (newest first), and when the next one will be."""
    supplier_id = _supplier_id(supplier_user)
    cutoff = approvals.approval_cutoff(db)
    last_closed = approvals._last_closed_delivery_date(cutoff)
    dates = [
        d
        for (d,) in db.query(SupplierPrice.delivery_date)
        .filter(SupplierPrice.supplier_id == supplier_id, SupplierPrice.delivery_date <= last_closed)
        .distinct()
        .order_by(SupplierPrice.delivery_date.desc())
        .limit(MAX_DATES)
    ]
    upcoming = (
        db.query(SupplierPrice.delivery_date)
        .filter(SupplierPrice.supplier_id == supplier_id, SupplierPrice.delivery_date > last_closed)
        .order_by(SupplierPrice.delivery_date)
        .first()
    )
    return {
        "dates": dates,
        "next_delivery_date": upcoming[0] if upcoming else None,
        "next_final_at": approvals.expires_at(upcoming[0], cutoff) if upcoming else None,
        "cutoff_label": approvals.format_cutoff(cutoff),
        "signature_ready": get_signature(db)["signature_image"] is not None,
    }


def _line(row: SupplierPrice, status: str | None, counter: float | None) -> tuple[float, str, float | None]:
    """(final price, basis, SPAR's adjusted price as the supplier saw it) for one quote."""
    own = float(row.price)
    sent = float(row.adjusted_price) if status not in (None, approvals.DRAFT) else None
    if status == approvals.APPROVED:
        return float(row.adjusted_price), BASIS_APPROVED, sent
    if status == approvals.REJECTED:
        return (counter, BASIS_COUNTER, sent) if counter is not None else (own, BASIS_REJECTED, sent)
    if status == approvals.EXPIRED:
        return own, BASIS_EXPIRED, sent
    return own, BASIS_OWN, None


def build_list(db: Session, supplier_user: User, delivery_date: date) -> dict:
    supplier_id = _supplier_id(supplier_user)
    cutoff = approvals.approval_cutoff(db)
    final_at = approvals.expires_at(delivery_date, cutoff)
    if approvals._now() < final_at:
        raise ValidationFailedError(
            f"Prices for {_submission_day(delivery_date):%A %d %B} become final at "
            f"{approvals.format_cutoff(cutoff)} on {final_at:%A %d %B}."
        )
    rows = (
        db.query(SupplierPrice)
        .filter(SupplierPrice.supplier_id == supplier_id, SupplierPrice.delivery_date == delivery_date)
        .all()
    )
    if not rows:
        raise NotFoundError(f"You didn't submit prices for {_submission_day(delivery_date):%A %d %B}.")

    products = {p.id: p for p in db.query(Product).filter(Product.id.in_([r.product_id for r in rows]))}
    categories = {c.id: c.name for c in db.query(ProductCategory)}
    statuses = approvals.row_statuses(db, rows)
    lines = []
    for r in rows:
        status, _, counter = statuses[r.id]
        final, basis, adjusted = _line(r, status, counter)
        p = products.get(r.product_id)
        lines.append(
            {
                "product_id": r.product_id,
                "product_code": p.product_code if p else "—",
                "category": categories.get(p.category_id, "") if p else "",
                "category_id": p.category_id if p else 0,
                "product_description": p.description if p else f"Product #{r.product_id}",
                "unit_code": r.unit_code,
                "your_price": float(r.price),
                "adjusted_price": adjusted,
                "final_price": final,
                "basis": basis,
            }
        )
    lines.sort(key=lambda l: (l["category_id"], l["product_description"]))

    signed = (
        db.query(SupplierPriceRevision)
        .filter(
            SupplierPriceRevision.supplier_id == supplier_id,
            SupplierPriceRevision.delivery_date == delivery_date,
            SupplierPriceRevision.status.in_((approvals.APPROVED, approvals.PARTIAL)),
            SupplierPriceRevision.signature_image.isnot(None),
        )
        .order_by(SupplierPriceRevision.responded_at)
        .all()
    )
    signers = {u.id: u.username for u in db.query(User).filter(User.id.in_({s.responded_by for s in signed if s.responded_by}))}
    supplier = db.get(Supplier, supplier_id)
    # A short fingerprint of the figures, printed on both files so a copy
    # can be matched back to what the system holds.
    reference = hashlib.sha256(
        json.dumps(
            [supplier_id, delivery_date.isoformat(), [(l["product_id"], l["final_price"]) for l in lines]],
            separators=(",", ":"),
        ).encode()
    ).hexdigest()[:12].upper()
    return {
        "supplier_name": supplier.supplier_name if supplier else "—",
        "delivery_date": delivery_date,
        "finalized_at": final_at,
        "reference": reference,
        "lines": lines,
        "spar_signature": get_signature(db),
        "supplier_signatures": [
            {
                "price_sheet_id": s.id,
                "signer_name": s.signer_name or signers.get(s.responded_by),
                "signed_at": s.responded_at.astimezone(approvals.BUSINESS_TZ) if s.responded_at else None,
                "signature_image": s.signature_image,
            }
            for s in signed
        ],
    }


def _require_spar_signature(data: dict) -> None:
    if not data["spar_signature"]["signature_image"]:
        raise ValidationFailedError(
            "SPAR hasn't added its signature to final price lists yet — the download will be available once it has."
        )


def _png_bytes(data_url: str) -> bytes:
    return base64.b64decode(data_url.split(",", 1)[1])


def _submission_day(delivery_date: date) -> date:
    """The day the supplier submitted these prices — shown to suppliers as "Date"."""
    return delivery_date - timedelta(days=approvals.SUBMISSION_DAYS_BEFORE_DELIVERY)


def _when(dt: datetime) -> str:
    return f"{approvals.format_cutoff(dt.time())} on {dt:%A %d %B %Y}"


def filename(data: dict, ext: str) -> str:
    safe = "".join(c if c.isalnum() else "-" for c in data["supplier_name"]).strip("-").lower() or "supplier"
    return f"spar-final-prices-{safe}-{_submission_day(data['delivery_date']).isoformat()}.{ext}"


# ---- Excel ----


def build_excel(data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    _require_spar_signature(data)
    headers = ["Product Code", "Category", "Product", "Unit", "Your Price", "SPAR Adjusted Price", "Final Price", "Basis"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Final Prices"
    last_col = len(headers)

    ws.row_dimensions[1].height = 42
    logo = XLImage(str(_SPAR_LOGO_PATH))
    logo.width, logo.height = 280, 44  # source is 1014x161
    ws.add_image(logo, "A1")

    def banner(row: int, text: str, font: Font, height: float) -> None:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        ws.cell(row=row, column=1, value=text).font = font
        ws.row_dimensions[row].height = height

    banner(2, f"Final Price List — {data['supplier_name']}", Font(bold=True, size=14, color=_SPAR_GREEN), 22)
    banner(3, f"Date: {_submission_day(data['delivery_date']):%A %d %B %Y}", Font(bold=True, size=11), 18)
    banner(
        4,
        f"Prices final at {_when(data['finalized_at'])}  ·  Reference {data['reference']}",
        Font(italic=True, size=9, color="666666"),
        16,
    )

    header_row = 6
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=_SPAR_GREEN, end_color=_SPAR_GREEN, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30

    final_fill = PatternFill(start_color="E8F5EC", end_color="E8F5EC", fill_type="solid")
    for l in data["lines"]:
        code = l["product_code"]
        code = int(code) if code.isdigit() and not code.startswith("0") else code
        ws.append(
            [code, l["category"], l["product_description"], l["unit_code"], l["your_price"], l["adjusted_price"], l["final_price"], l["basis"]]
        )
        r = ws.max_row
        ws.cell(row=r, column=1).number_format = "0" if isinstance(code, int) else "@"
        for col in (5, 6, 7):
            ws.cell(row=r, column=col).number_format = "#,##0.00"
        ws.cell(row=r, column=7).font = Font(bold=True)
        ws.cell(row=r, column=7).fill = final_fill

    for col, width in zip("ABCDEFGH", (14, 12, 36, 7, 12, 14, 12, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = f"A{header_row + 1}"

    # Signatures below the table: SPAR on the left, the supplier on the right.
    row = ws.max_row + 3
    rule = Border(top=Side(style="thin", color="999999"))

    def signature_block(col: int, heading: str, image: str | None, lines: list[str]) -> None:
        ws.cell(row=row, column=col, value=heading).font = Font(bold=True, color=_SPAR_GREEN)
        if image:
            pic = XLImage(io.BytesIO(_png_bytes(image)))
            scale = min(200 / pic.width, 70 / pic.height, 1)
            pic.width, pic.height = pic.width * scale, pic.height * scale
            ws.add_image(pic, ws.cell(row=row + 1, column=col).coordinate)
        for i, text in enumerate(lines):
            cell = ws.cell(row=row + 5 + i, column=col, value=text)
            cell.font = Font(size=9, bold=i == 0, color="333333" if i == 0 else "666666")
            if i == 0:
                cell.border = rule

    for r_ in range(row + 1, row + 5):
        ws.row_dimensions[r_].height = 18
    spar = data["spar_signature"]
    signature_block(
        3,
        "Authorised for SPAR",
        spar["signature_image"],
        [spar["signatory_name"] or "", spar["signatory_title"] or "", f"Signed electronically, {_when(data['finalized_at'])}"],
    )
    sup = data["supplier_signatures"][-1] if data["supplier_signatures"] else None
    if sup:
        signature_block(
            6,
            f"Accepted by {data['supplier_name']}",
            sup["signature_image"],
            [sup["signer_name"] or "", f"E-signed price sheet #{sup['price_sheet_id']}", f"{sup['signed_at']:%d %B %Y %H:%M}"],
        )
    else:
        ws.cell(row=row, column=6, value="Supplier").font = Font(bold=True, color=_SPAR_GREEN)
        ws.cell(row=row + 1, column=6, value="No SPAR adjustment was approved for this date — no supplier signature needed.").font = Font(
            italic=True, size=9, color="666666"
        )

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- PDF ----


def build_pdf(data: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    _require_spar_signature(data)
    green = colors.HexColor("#" + _SPAR_GREEN)
    grey = colors.HexColor("#666666")
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=8.5, leading=10.5)
    small = ParagraphStyle("small", parent=body, fontSize=7.5, leading=9.5, textColor=grey)
    title = ParagraphStyle("title", parent=styles["Title"], fontSize=16, leading=20, textColor=green, alignment=0, spaceAfter=2)
    heading = ParagraphStyle("heading", parent=body, fontName="Helvetica-Bold", textColor=green, fontSize=9.5)
    final_style = ParagraphStyle("final", parent=body, fontName="Helvetica-Bold", alignment=TA_RIGHT)

    def esc(s: str) -> str:
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def money(v: float | None) -> str:
        return f"{v:,.2f}" if v is not None else "—"

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
        title=f"SPAR final price list — {data['supplier_name']} — {_submission_day(data['delivery_date']).isoformat()}",
        author="SPAR Fresh Procurement",
    )
    story = [
        Image(str(_SPAR_LOGO_PATH), width=60 * mm, height=60 * mm * 161 / 1014),
        Spacer(1, 4 * mm),
        Paragraph(f"Final Price List — {esc(data['supplier_name'])}", title),
        Paragraph(f"<b>Date:</b> {_submission_day(data['delivery_date']):%A %d %B %Y}", body),
        Paragraph(f"Prices final at {_when(data['finalized_at'])} &nbsp;·&nbsp; Reference {data['reference']}", small),
        Spacer(1, 5 * mm),
    ]

    header = ["Code", "Product", "Unit", "Your Price", "SPAR Adjusted", "Final Price", "Basis"]
    table_rows = [header]
    for l in data["lines"]:
        table_rows.append(
            [
                l["product_code"],
                Paragraph(esc(l["product_description"]), body),
                l["unit_code"],
                money(l["your_price"]),
                money(l["adjusted_price"]),
                Paragraph(money(l["final_price"]), final_style),
                Paragraph(esc(l["basis"]), small),
            ]
        )
    table = Table(table_rows, colWidths=[16 * mm, 53 * mm, 11 * mm, 21 * mm, 26 * mm, 21 * mm, 34 * mm], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), green),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("ALIGN", (3, 0), (5, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F8F5")]),
                ("BACKGROUND", (5, 1), (5, -1), colors.HexColor("#E8F5EC")),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5E2D9")),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story += [table, Spacer(1, 10 * mm)]

    def signature_cell(heading_text: str, image: str | None, lines: list[str], placeholder: str | None = None) -> list:
        cell = [Paragraph(esc(heading_text), heading), Spacer(1, 2 * mm)]
        if image:
            img = Image(io.BytesIO(_png_bytes(image)))
            scale = min(60 * mm / img.imageWidth, 22 * mm / img.imageHeight)
            img.drawWidth, img.drawHeight = img.imageWidth * scale, img.imageHeight * scale
            img.hAlign = "LEFT"
            cell.append(img)
        else:
            cell += [Spacer(1, 4 * mm), Paragraph(esc(placeholder or ""), small), Spacer(1, 4 * mm)]
        cell.append(Table([[""]], colWidths=[70 * mm], style=[("LINEABOVE", (0, 0), (-1, -1), 0.6, grey)]))
        for i, text in enumerate(t for t in lines if t):
            cell.append(Paragraph(f"<b>{esc(text)}</b>" if i == 0 else esc(text), body if i == 0 else small))
        return cell

    spar = data["spar_signature"]
    left = signature_cell(
        "Authorised for SPAR",
        spar["signature_image"],
        [spar["signatory_name"] or "", spar["signatory_title"] or "", f"Signed electronically, {_when(data['finalized_at'])}"],
    )
    sup = data["supplier_signatures"][-1] if data["supplier_signatures"] else None
    right_cell = (
        signature_cell(
            f"Accepted by {data['supplier_name']}",
            sup["signature_image"],
            [sup["signer_name"] or "", f"E-signed price sheet #{sup['price_sheet_id']}, {sup['signed_at']:%d %B %Y %H:%M}"],
        )
        if sup
        else signature_cell(
            "Supplier",
            None,
            [],
            placeholder="No SPAR adjustment was approved for this date, so no supplier signature was needed.",
        )
    )
    signatures = Table([[left, right_cell]], colWidths=[91 * mm, 91 * mm])
    signatures.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story.append(KeepTogether([signatures]))

    footer_text = f"SPAR final price list · {data['supplier_name']} · date {_submission_day(data['delivery_date']).isoformat()} · Ref {data['reference']}"

    def footer(canvas, doc_) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(grey)
        canvas.drawString(14 * mm, 9 * mm, footer_text)
        canvas.drawRightString(A4[0] - 14 * mm, 9 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
