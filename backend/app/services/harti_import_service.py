"""
Daily local market price import from the HARTI (Hector Kobbekaduwa
Agrarian Research and Training Institute) wholesale vegetable bulletin.

This is a port of the standalone market-price-system tool
(downloadLatestHartiPdf.ts, readPdf.ts, calculations.ts,
productMapping.ts) into this backend's own Python stack, so results
match the Excel that tool used to produce:

1. Read HARTI's daily-price index page — one table row per bulletin
   (date | medium | PDF link) — and pick the newest *English* bulletin.
   The page is plain server-rendered HTML, so no headless browser is
   needed (the standalone tool used Playwright only to read that table).
2. Read the bulletin's wholesale market table by text position: each
   market column is located from its heading word, and every word on a
   product row is assigned to the column it starts in.
3. For the mapped products only (PRODUCT_MAPPING — PDF label → DC code
   + system name), take the min–max range in four markets (Dambulla,
   Thambuththegama, Keppetipola, Nuwara Eliya), each market's midpoint,
   and the final average of whichever of those markets have a price.
4. REPLACE that report date's rows in local_market_prices, and write
   each final average as the LOCAL_MARKET reference price for the same
   date, so the Price History page keeps a Local Market trend.

Runs three ways:
- scripts/import_harti_prices.py — the scheduled daily worker.
- POST /local-market-prices/sync — admin's "Fetch latest from HARTI".
- POST /local-market-prices/upload — admin uploads a bulletin PDF by
  hand, for when harti.gov.lk is down or unreachable from the server.
"""
import html
import io
import re
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import NamedTuple

import openpyxl
import pdfplumber
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models.local_market_price import LocalMarketPrice
from app.models.market_reference_price import MarketReferencePrice
from app.models.product import Product
from app.models.user import Role, User, UserRole
from app.schemas._limits import MAX_NUMERIC_10_2

SOURCE = "LOCAL_MARKET"
REQUEST_TIMEOUT = 60
# A HARTI bulletin is well under 1 MB; this only rejects something
# clearly wrong before pdfplumber ever opens it.
MAX_PDF_BYTES = 20 * 1024 * 1024
_HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; SPAR-procurement-price-import/1.0)"}

HARTI_INDEX_URL = "https://www.harti.gov.lk/daily-price.php"


class HartiImportError(Exception):
    """A bulletin couldn't be fetched, read, or matched — the caller (the
    scheduled worker or an API route) logs/surfaces this and moves on.
    Raised before anything is written, so existing prices stay untouched."""


# =============================================================
# PRODUCT MAPPING — PDF label → DC code + system name.
# Ported as-is from the standalone tool's productMapping.ts. A None
# value explicitly excludes that PDF product from the report.
# =============================================================


class MappedProduct(NamedTuple):
    dc_code: str
    system_name: str


PRODUCT_MAPPING: dict[str, MappedProduct | None] = {
    "Beans": MappedProduct("4503951", "GREEN BEANS"),
    "Carrot": MappedProduct("4503950", "CARROTS"),
    "Leeks": MappedProduct("4503953", "LEEKS"),
    "Beet root": MappedProduct("4512940", "CUT BEETROOT"),
    "Beet root (N Eliya)": MappedProduct("4503947", "BEETROOT"),
    "Knolkhol": MappedProduct("4503952", "KNOL KHOL"),
    "Raddish": MappedProduct("4503955", "RADDISH"),
    "Cabbage (N'Eliya)": MappedProduct("4503948", "CABBAGE"),
    "Cabbage (Kandy)": None,
    "Tomato": MappedProduct("4504011", "TOMATOES"),
    "Ladies Fingers": MappedProduct("4504026", "LADIES FINGERS"),
    "Brinjals": MappedProduct("4504016", "BRINJALS"),
    "Capsicum": MappedProduct("4503949", "CAPSICUM"),
    "Pumpkin": MappedProduct("4504033", "PUMPKIN"),
    "Cucumber": MappedProduct("4504017", "CUCUMBER"),
    "Bitter Gourd": MappedProduct("4504015", "BITTER GOURD"),
    "Snake Gourd": MappedProduct("4504035", "SNAKE GOURD"),
    "Drumstick": MappedProduct("4504020", "DRUMSTICKS"),
    "Luffa": MappedProduct("4504034", "RIBBED GOURD"),
    "Long Beans": MappedProduct("4503767", "LONG BEANS"),
    "Ash Plantains": MappedProduct("4504012", "ASH PLANTAINS"),
    "Green Chillies": MappedProduct("4503764", "GREEN CHILIES"),
    "Lime": MappedProduct("4504029", "LIME"),
    "Sweet Potatoe": MappedProduct("4503772", "SWEET POTATO"),
    "Manioc": MappedProduct("4503604", "MANIOC"),
    "Eggplant": MappedProduct("4503763", "EGG PLANTS"),
    "Potato(Imported)": None,
    "Potato (Imported)": None,
    "Potato (Welimada)": None,
    "Potato (Nuwaraeliya)": MappedProduct("4503954", "LOCAL POTATOES"),
    "B'Onion Imported": None,
    "Big-onion Local": MappedProduct("4503211", "BIG ONIONS"),
}


# =============================================================
# MARKETS
# =============================================================


@dataclass(frozen=True)
class Market:
    key: str
    name: str
    aliases: tuple[str, ...]


# The four markets reported on, in display order. `key` doubles as the
# column-name prefix on LocalMarketPrice (e.g. dambulla_min).
MARKETS: tuple[Market, ...] = (
    Market("dambulla", "Dambulla", ("Dambulla",)),
    Market("thambuththegama", "Thambuththegama", ("Thambuththegama", "Thambuththe gama")),
    Market("keppetipola", "Keppetipola", ("Keppetipola",)),
    Market("nuwara_eliya", "Nuwara Eliya", ("Nuwara Eliya", "Nuwaraeliya", "N'Eliya")),
)

# Every market heading in the bulletin, reported on or not — column
# boundaries are the midpoints between neighbouring headings, so the
# unreported columns are needed to fence in the four that are.
ALL_MARKET_HEADINGS: tuple[str, ...] = (
    "Peliyagoda",
    "Kandy",
    "Dambulla",
    "Meegoda",
    "Norochchole",
    "Thambuththegama",
    "Keppetipola",
    "Nuwara Eliya",
    "Bandarawela",
    "Veyangoda",
)


# =============================================================
# PARSED RESULT
# =============================================================

_CENT = Decimal("0.01")


def _round2(value: Decimal) -> Decimal:
    # Half-up, matching the standalone tool's Number(x.toFixed(2)) — Python's
    # round() is half-even and would differ on e.g. 313.625.
    return value.quantize(_CENT, rounding=ROUND_HALF_UP)


@dataclass(frozen=True)
class PriceRange:
    min: Decimal
    max: Decimal

    @property
    def average(self) -> Decimal:
        """The market's midpoint — calculateMarketAverages in the standalone tool."""
        return _round2((self.min + self.max) / 2)


def final_average(ranges: list[PriceRange]) -> Decimal | None:
    """Mean of the available market midpoints; a market with no price is
    left out rather than counted as zero — calculateFinalAverage."""
    if not ranges:
        return None
    return _round2(sum((r.average for r in ranges), Decimal(0)) / len(ranges))


@dataclass
class BulletinRow:
    dc_code: str
    system_name: str
    pdf_name: str
    markets: dict[str, PriceRange]  # market key → range; only markets with a price

    @property
    def final_average(self) -> Decimal:
        return final_average(list(self.markets.values()))


@dataclass
class Bulletin:
    report_date: date
    rows: list[BulletinRow]
    # Mapped products listed in the bulletin but with no price in any of
    # the four markets that day — left out of `rows`, reported back to Admin.
    not_priced: list[str] = field(default_factory=list)


# =============================================================
# TEXT HELPERS
# =============================================================


def _normalize(value: str) -> str:
    value = unicodedata.normalize("NFD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]", "", value.lower())


_MAPPING_BY_KEY = {_normalize(label): mapped for label, mapped in PRODUCT_MAPPING.items()}

# Digit lookarounds rather than \b, which treats "_" as a word character
# and so missed the date in a name like "harti_2026-10-02.pdf".
_DATE_RE = re.compile(r"(?<!\d)(20\d{2})[./_-](\d{1,2})[./_-](\d{1,2})(?!\d)")


def _valid_date(year: str, month: str, day: str) -> date | None:
    try:
        return date(int(year), int(month), int(day))
    except ValueError:
        return None


def extract_date(text: str) -> date | None:
    """First valid YYYY.MM.DD / YYYY-MM-DD / YYYY_MM_DD date in `text`."""
    for match in _DATE_RE.finditer(text):
        parsed = _valid_date(*match.groups())
        if parsed:
            return parsed
    return None


_RANGE_RE = re.compile(r"(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)")
_SINGLE_PRICE_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*$")
_MAX_PRICE = Decimal(str(MAX_NUMERIC_10_2))


def parse_price_range(text: str) -> PriceRange | None:
    """'450- 500' → 450..500; a lone '450' → 450..450; '-' or blank → None."""
    cleaned = re.sub(r"\s+", " ", re.sub(r"[–—−]", "-", text.replace(",", ""))).strip()
    if not cleaned or cleaned == "-":
        return None
    try:
        match = _RANGE_RE.search(cleaned)
        if match:
            low, high = Decimal(match.group(1)), Decimal(match.group(2))
        else:
            single = _SINGLE_PRICE_RE.match(cleaned)
            if not single:
                return None
            low = high = Decimal(single.group(1))
    except InvalidOperation:
        return None
    if low > high or high > _MAX_PRICE:
        return None
    return PriceRange(min=low, max=high)


# =============================================================
# PDF LAYOUT — words, rows, columns
# =============================================================


@dataclass(frozen=True)
class _Word:
    text: str
    x0: float
    x1: float
    y: float  # baseline-ish, measured down from the top of the page

    @property
    def center(self) -> float:
        return (self.x0 + self.x1) / 2


@dataclass(frozen=True)
class _Column:
    key: str
    left: float
    right: float


def _page_words(page) -> list[_Word]:
    return [
        _Word(text=w["text"].strip(), x0=float(w["x0"]), x1=float(w["x1"]), y=float(w["bottom"]))
        for w in page.extract_words()
        if w["text"].strip()
    ]


def _group_rows(words: list[_Word], tolerance: float = 2.5) -> list[list[_Word]]:
    """Words whose baselines are within `tolerance` of a row's first word share that row."""
    rows: list[tuple[float, list[_Word]]] = []
    for word in sorted(words, key=lambda w: w.y):
        for anchor_y, row in rows:
            if abs(anchor_y - word.y) <= tolerance:
                row.append(word)
                break
        else:
            rows.append((word.y, [word]))
    return [sorted(row, key=lambda w: w.x0) for _, row in rows]


def _adjacent_pairs(rows: list[list[_Word]], max_gap: float = 6.0) -> list[_Word]:
    """
    Two neighbouring words on a line, joined — pdfplumber splits text on
    spaces, so a two-word heading ("Nuwara Eliya") only exists as a pair.
    """
    pairs = []
    for row in rows:
        for a, b in zip(row, row[1:]):
            if b.x0 - a.x1 <= max_gap:
                pairs.append(_Word(text=f"{a.text} {b.text}", x0=a.x0, x1=b.x1, y=a.y))
    return pairs


def _find_heading(words: list[_Word], pairs: list[_Word], aliases: tuple[str, ...]) -> _Word | None:
    """Topmost word matching an alias (exactly, or containing a 5+ char alias); joined pairs must match exactly."""
    norm_aliases = [_normalize(a) for a in aliases]
    candidates = [
        w
        for w in words
        if any(_normalize(w.text) == a or (len(a) >= 5 and a in _normalize(w.text)) for a in norm_aliases)
    ]
    candidates += [p for p in pairs if _normalize(p.text) in norm_aliases]
    return min(candidates, key=lambda w: w.y) if candidates else None


def _detect_headings(
    words: list[_Word], pairs: list[_Word], names: list[tuple[str, tuple[str, ...]]]
) -> list[tuple[str, float]]:
    """(name, center x) for every heading in `names` found on the page, left to right."""
    found = []
    for name, aliases in names:
        heading = _find_heading(words, pairs, aliases)
        if heading:
            found.append((name, heading.center))
    return sorted(found, key=lambda h: h[1])


def _build_columns(reported: list[tuple[str, float]], all_headings: list[tuple[str, float]]) -> list[_Column]:
    """Each reported market's column spans midway to its neighbouring headings on either side."""
    positions = all_headings if len(all_headings) >= 2 else reported
    index_by_name = {_normalize(name): i for i, (name, _) in enumerate(positions)}
    key_by_name = {_normalize(m.name): m.key for m in MARKETS}
    columns = []
    for name, _ in reported:
        i = index_by_name.get(_normalize(name))
        if i is None:
            continue
        center = positions[i][1]
        prev_center = positions[i - 1][1] if i > 0 else None
        next_center = positions[i + 1][1] if i + 1 < len(positions) else None
        if prev_center is not None:
            left = (prev_center + center) / 2
        else:
            left = center - ((next_center - center) / 2 if next_center is not None else 25)
        if next_center is not None:
            right = (center + next_center) / 2
        else:
            right = center + ((center - prev_center) / 2 if prev_center is not None else 25)
        columns.append(_Column(key=key_by_name[_normalize(name)], left=left, right=right))
    return columns


def _product_column_edge(all_headings: list[tuple[str, float]]) -> float | None:
    """Everything left of the first market column is the product name."""
    if len(all_headings) < 2:
        return None
    first, second = all_headings[0][1], all_headings[1][1]
    return first - (second - first) / 2


def _cell_text(row: list[_Word], column: _Column) -> str:
    return " ".join(w.text for w in row if column.left <= w.x0 < column.right).strip()


# =============================================================
# BULLETIN PARSING
# =============================================================

_REPORTED_MARKETS = [(m.name, m.aliases) for m in MARKETS]
_ALL_HEADINGS = [(name, (name,)) for name in ALL_MARKET_HEADINGS]


def parse_bulletin(pdf_bytes: bytes, report_date: date | None = None, filename: str | None = None) -> Bulletin:
    """
    Reads every mapped product's four-market ranges out of a HARTI
    bulletin PDF. The report date is, in order: `report_date` if given
    (the date on HARTI's own index page), a date in `filename`, or the
    first date printed in the PDF itself.

    Market headings are looked for on every page; a page without its own
    headings reuses the last page's columns, in case the table ever runs
    over a page break. The FIRST row found for a product is final — later
    pages hold the Sinhala copy of the table and summary charts whose
    labels ("Capsicum (% ↑)") can normalize to a product name, and must
    never fill in a product the real table left blank.
    """
    if report_date is None and filename:
        report_date = extract_date(filename)

    pages: list[list[_Word]] = []
    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            for page in pdf.pages:
                try:
                    pages.append(_page_words(page))
                except Exception:
                    pages.append([])  # an unreadable page (scanned image, broken font) — the table is elsewhere
    except Exception as e:  # pdfplumber/pdfminer raise a wide variety of errors on a bad file
        raise HartiImportError(f"Could not read this PDF ({e.__class__.__name__}) — is it a HARTI daily price bulletin?")
    return _parse_pages(pages, report_date)


def _parse_pages(pages: list[list[_Word]], report_date: date | None) -> Bulletin:
    rows: list[BulletinRow] = []
    not_priced: list[str] = []
    seen: set[str] = set()
    reported: list[tuple[str, float]] = []
    all_headings: list[tuple[str, float]] = []

    for words in pages:
        if not words:
            continue
        if report_date is None:
            report_date = extract_date(" ".join(w.text for w in words))

        page_rows = _group_rows(words)
        pairs = _adjacent_pairs(page_rows)
        page_reported = _detect_headings(words, pairs, _REPORTED_MARKETS)
        page_all = _detect_headings(words, pairs, _ALL_HEADINGS)
        if len(page_reported) >= 2:
            reported = page_reported
        if len(page_all) >= 2:
            all_headings = page_all
        if len(reported) < 2 or len(all_headings) < 2:
            continue

        columns = _build_columns(reported, all_headings)
        name_edge = _product_column_edge(all_headings)
        if name_edge is None or len(columns) < 2:
            continue

        for row in page_rows:
            pdf_name = " ".join(w.text for w in row if w.x0 < name_edge).strip()
            if not pdf_name:
                continue
            mapped = _MAPPING_BY_KEY.get(_normalize(pdf_name))
            if mapped is None or mapped.dc_code in seen:
                continue
            seen.add(mapped.dc_code)

            markets = {}
            for column in columns:
                price_range = parse_price_range(_cell_text(row, column))
                if price_range:
                    markets[column.key] = price_range
            if not markets:
                not_priced.append(mapped.system_name)
                continue
            rows.append(
                BulletinRow(dc_code=mapped.dc_code, system_name=mapped.system_name, pdf_name=pdf_name, markets=markets)
            )

    if report_date is None:
        raise HartiImportError("Could not find the report date in this PDF or its filename.")
    if not rows:
        raise HartiImportError(
            "No mapped product prices were found in this PDF — check that it is the English HARTI daily "
            "price bulletin (the wholesale market table may have changed layout)."
        )
    return Bulletin(report_date=report_date, rows=rows, not_priced=not_priced)


# =============================================================
# FETCHING THE LATEST BULLETIN FROM HARTI
# =============================================================

_ROW_RE = re.compile(r"<tr\b[^>]*>(.*?)</tr>", re.IGNORECASE | re.DOTALL)
_TD_RE = re.compile(r"<td\b[^>]*>(.*?)</td>", re.IGNORECASE | re.DOTALL)
_HREF_RE = re.compile(r"""<a\b[^>]*\bhref\s*=\s*["']([^"']+)["']""", re.IGNORECASE)
_TAG_RE = re.compile(r"<[^>]+>")
_PDF_HREF_RE = re.compile(r"\.pdf(?:$|[?#])", re.IGNORECASE)


def _get(url: str) -> bytes:
    req = urllib.request.Request(url, headers=_HEADERS)
    try:
        with urllib.request.urlopen(req, timeout=REQUEST_TIMEOUT) as resp:  # noqa: S310 — https only, see callers
            return resp.read(MAX_PDF_BYTES + 1)
    except urllib.error.HTTPError as e:
        raise HartiImportError(f"harti.gov.lk returned HTTP {e.code} for {url}.")
    except (urllib.error.URLError, TimeoutError, OSError) as e:
        reason = getattr(e, "reason", e)
        raise HartiImportError(f"Could not reach harti.gov.lk ({reason}). Try again later, or upload the PDF instead.")


def _cell_plain_text(cell_html: str) -> str:
    return html.unescape(_TAG_RE.sub(" ", cell_html)).strip()


def find_latest_bulletin() -> tuple[date, str]:
    """(report date, absolute PDF URL) of the newest English bulletin on HARTI's index page."""
    page = _get(HARTI_INDEX_URL).decode("utf-8", errors="ignore")

    english: list[tuple[date, str]] = []
    for row_html in _ROW_RE.findall(page):
        cells = _TD_RE.findall(row_html)
        if len(cells) < 3:
            continue
        date_text, medium = _cell_plain_text(cells[0]), _cell_plain_text(cells[1])
        href_match = _HREF_RE.search(cells[2])
        if not href_match or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", date_text):
            continue
        # The parser reads English product and market names — never fall
        # back to a Sinhala or Tamil bulletin.
        if "english" not in medium.lower():
            continue
        href = html.unescape(href_match.group(1)).strip()
        report_date = _valid_date(*date_text.split("-"))
        if report_date and _PDF_HREF_RE.search(href):
            english.append((report_date, href))

    if not english:
        raise HartiImportError(
            "No English PDF bulletins were found on the HARTI daily price page — its layout may have changed."
        )

    report_date, href = max(english, key=lambda r: r[0])
    # Bulletin filenames contain spaces and parentheses, e.g.
    # "Vegetable Pricenew ex1(2026.10.02).pdf".
    url = urllib.parse.urljoin(HARTI_INDEX_URL, urllib.parse.quote(href, safe="/:?#&=%()"))
    parts = urllib.parse.urlsplit(url)
    if parts.scheme != "https" or not (parts.hostname or "").endswith("harti.gov.lk"):
        raise HartiImportError(f"Refusing to download a bulletin from outside harti.gov.lk: {url}")
    return report_date, url


def _check_pdf(pdf_bytes: bytes) -> None:
    if len(pdf_bytes) > MAX_PDF_BYTES:
        raise HartiImportError("That file is too large to be a HARTI daily price bulletin.")
    if not pdf_bytes.startswith(b"%PDF-"):
        raise HartiImportError("That file is not a PDF.")


def fetch_latest_bulletin() -> Bulletin:
    report_date, url = find_latest_bulletin()
    pdf_bytes = _get(url)
    _check_pdf(pdf_bytes)
    return parse_bulletin(pdf_bytes, report_date=report_date)


# =============================================================
# SAVING / READING
# =============================================================


def _find_an_admin(db: Session) -> User | None:
    return (
        db.query(User)
        .join(UserRole, UserRole.user_id == User.id)
        .join(Role, Role.id == UserRole.role_id)
        .filter(Role.code == "ADMIN", User.is_active.is_(True))
        .order_by(User.id)
        .first()
    )


def save_bulletin(db: Session, admin: User, bulletin: Bulletin) -> dict:
    """
    REPLACES everything stored for the bulletin's report date — its
    local_market_prices rows and its LOCAL_MARKET reference prices — in
    one transaction, so a re-import never leaves a stale product behind
    and a failure never leaves the day half-written.
    """
    report_date = bulletin.report_date
    codes = [r.dc_code for r in bulletin.rows]
    products = {p.product_code: p for p in db.query(Product).filter(Product.product_code.in_(codes)).all()}
    now = datetime.now(timezone.utc)

    try:
        db.query(LocalMarketPrice).filter(LocalMarketPrice.report_date == report_date).delete(
            synchronize_session=False
        )
        db.query(MarketReferencePrice).filter(
            MarketReferencePrice.source == SOURCE, MarketReferencePrice.delivery_date == report_date
        ).delete(synchronize_session=False)

        unmatched = []
        for row in bulletin.rows:
            product = products.get(row.dc_code)
            columns = {}
            for market in MARKETS:
                price_range = row.markets.get(market.key)
                columns[f"{market.key}_min"] = price_range.min if price_range else None
                columns[f"{market.key}_max"] = price_range.max if price_range else None
            db.add(
                LocalMarketPrice(
                    report_date=report_date,
                    product_id=product.id if product else None,
                    dc_code=row.dc_code,
                    system_name=row.system_name,
                    pdf_name=row.pdf_name,
                    final_average=row.final_average,
                    imported_by=admin.id,
                    imported_at=now,
                    **columns,
                )
            )
            if product:
                db.add(
                    MarketReferencePrice(
                        product_id=product.id,
                        source=SOURCE,
                        delivery_date=report_date,
                        price=row.final_average,
                        updated_by=admin.id,
                        updated_at=now,
                    )
                )
            else:
                unmatched.append({"dc_code": row.dc_code, "system_name": row.system_name})
        db.commit()
    except Exception:
        db.rollback()
        raise

    return {
        "report_date": report_date,
        "saved": len(bulletin.rows),
        "not_priced": bulletin.not_priced,
        "unmatched": unmatched,
    }


def run_import(db: Session, admin: User | None = None) -> dict:
    """Fetches the newest English bulletin from HARTI and saves it (see save_bulletin)."""
    admin = admin or _find_an_admin(db)
    if not admin:
        raise HartiImportError("No active ADMIN account found to attribute this import to.")
    return save_bulletin(db, admin, fetch_latest_bulletin())


def import_pdf(db: Session, admin: User, pdf_bytes: bytes, filename: str | None) -> dict:
    """Saves a bulletin PDF Admin uploaded by hand (see save_bulletin)."""
    _check_pdf(pdf_bytes)
    return save_bulletin(db, admin, parse_bulletin(pdf_bytes, filename=filename))


def latest_report_dates(db: Session, limit: int = 90) -> list[date]:
    """Report dates that have imported prices, newest first."""
    return [
        d
        for (d,) in db.query(LocalMarketPrice.report_date)
        .distinct()
        .order_by(LocalMarketPrice.report_date.desc())
        .limit(limit)
        .all()
    ]


def get_report(db: Session, report_date: date | None = None) -> tuple[date | None, list[LocalMarketPrice]]:
    """One report date's rows in bulletin order — the newest report date if none is given."""
    if report_date is None:
        report_date = db.query(func.max(LocalMarketPrice.report_date)).scalar()
        if report_date is None:
            return None, []
    rows = (
        db.query(LocalMarketPrice)
        .filter(LocalMarketPrice.report_date == report_date)
        .order_by(LocalMarketPrice.id)
        .all()
    )
    return report_date, rows


def market_range(row: LocalMarketPrice, market_key: str) -> PriceRange | None:
    low, high = getattr(row, f"{market_key}_min"), getattr(row, f"{market_key}_max")
    if low is None or high is None:
        return None
    return PriceRange(min=Decimal(low), max=Decimal(high))


# =============================================================
# EXCEL EXPORT — same layout as the standalone tool's exportExcel.ts
# =============================================================


def build_excel(report_date: date, rows: list[LocalMarketPrice]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Market Prices"

    headers: list[tuple[str, int]] = [("Date", 15), ("DC Code", 15), ("System Product", 25), ("PDF Product", 25)]
    for market in MARKETS:
        width = len(market.name) + 10
        headers += [(f"{market.name} Min", width), (f"{market.name} Max", width), (f"{market.name} Average", width + 4)]
    headers.append(("Final Average", 18))
    sheet.append([h for h, _ in headers])

    for row in rows:
        values: list = [report_date.isoformat(), row.dc_code, row.system_name, row.pdf_name]
        for market in MARKETS:
            price_range = market_range(row, market.key)
            values += (
                [float(price_range.min), float(price_range.max), float(price_range.average)]
                if price_range
                else [None, None, None]
            )
        values.append(float(row.final_average))
        sheet.append(values)

    header_font = Font(bold=True, color="FFFFFFFF")
    header_fill = PatternFill("solid", fgColor="FF1F4E78")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for idx, (_, width) in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=idx)
        cell.font, cell.fill, cell.alignment = header_font, header_fill, header_align
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.row_dimensions[1].height = 32

    for column_cells in sheet.iter_cols(min_col=5, max_col=len(headers), min_row=2):
        for cell in column_cells:
            cell.number_format = "0.00"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
