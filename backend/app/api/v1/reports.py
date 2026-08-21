from datetime import date, timedelta
import io

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import get_current_user, require_roles
from app.core.errors import ValidationFailedError
from app.models.user import User
from app.schemas.report import AdminReportOut
from app.services import report_service

router = APIRouter(prefix="/reports", dependencies=[Depends(get_current_user)])


def _validate_range(start_date: date, end_date: date) -> None:
    if start_date > end_date:
        raise ValidationFailedError("start_date must not be after end_date.")
    if (end_date - start_date).days > 92:
        raise ValidationFailedError("Date range too large — please choose 92 days or fewer.")


@router.get("/admin/summary", response_model=AdminReportOut)
def get_report(
    start_date: date | None = None,
    end_date: date | None = None,
    admin: User = Depends(require_roles("ADMIN")),
    db: Session = Depends(get_db),
):
    """
    Defaults to the last 7 days (inclusive of today) if no range is given.
    Spend figures are based on Admin's agreed prices with suppliers, not
    branch demand — demand only becomes money once it's assigned.
    """
    if end_date is None:
        end_date = date.today()
    if start_date is None:
        start_date = end_date - timedelta(days=6)
    _validate_range(start_date, end_date)
    return report_service.build_report(db, start_date, end_date)


@router.get("/admin/summary/export")
def export_report(
    start_date: date | None = None,
    end_date: date | None = None,
    admin: User = Depends(require_roles("ADMIN")),
    db: Session = Depends(get_db),
):
    """Same data as /admin/summary, as a downloadable .xlsx workbook with one sheet per section."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    if end_date is None:
        end_date = date.today()
    if start_date is None:
        start_date = end_date - timedelta(days=6)
    _validate_range(start_date, end_date)
    report = report_service.build_report(db, start_date, end_date)

    wb = Workbook()
    header_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    def write_sheet(ws, headers, rows):
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for row in rows:
            ws.append(row)
        for i, _ in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = 20

    ws1 = wb.active
    ws1.title = "Daily Summary"
    write_sheet(
        ws1,
        ["Delivery Date", "Orders", "Assigned Orders", "Branches", "Spend (Rs.)"],
        [[d.delivery_date.isoformat(), d.orders_count, d.assigned_orders_count, d.branches_count, d.spend] for d in report.daily],
    )

    ws2 = wb.create_sheet("By Supplier")
    write_sheet(
        ws2,
        ["Supplier Code", "Supplier", "Spend (Rs.)", "Assignment Lines"],
        [[s.supplier_code, s.supplier_name, s.spend, s.lines_count] for s in report.by_supplier],
    )

    ws3 = wb.create_sheet("By Category")
    write_sheet(
        ws3,
        ["Category", "Spend (Rs.)", "Assignment Lines"],
        [[c.category_name, c.spend, c.lines_count] for c in report.by_category],
    )

    ws4 = wb.create_sheet("Top Products")
    write_sheet(
        ws4,
        ["Code", "Product", "Category", "Unit", "Total Quantity", "Total Spend (Rs.)", "Orders"],
        [
            [p.product_code, p.description, p.category_name, p.unit_code, p.total_quantity, p.total_spend, p.order_count]
            for p in report.top_products
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"report-{start_date.isoformat()}-to-{end_date.isoformat()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
