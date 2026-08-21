from datetime import date
import io

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import get_current_user, require_roles
from app.models.user import User
from app.schemas.supplier_order import (
    SetSupplierOrderRequest,
    SupplierOrderAdminOut,
    MySupplierOrdersOut,
    SupplierPricePreviewOut,
)
from app.services import supplier_order_service

router = APIRouter(prefix="/supplier-orders", dependencies=[Depends(get_current_user)])


@router.get("/mine", response_model=MySupplierOrdersOut)
def my_supplier_orders(
    delivery_date: date | None = None,
    current_user: User = Depends(require_roles("SUPPLIER")),
    db: Session = Depends(get_db),
):
    """A supplier's own orders, grouped branch by branch, for one delivery date."""
    return supplier_order_service.get_my_orders_by_branch(db, current_user, delivery_date)


@router.get("/mine/export")
def export_my_supplier_orders(
    delivery_date: date | None = None,
    current_user: User = Depends(require_roles("SUPPLIER")),
    db: Session = Depends(get_db),
):
    """Same data as /mine, pivoted into product-rows x branch-columns and
    downloaded as .xlsx — mirrors the on-screen Orders by Branch table."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    data = supplier_order_service.get_my_orders_by_branch(db, current_user, delivery_date)

    # Pivot: same transformation the frontend does — one row per product,
    # one column per branch, quantity in the cell.
    branches = [{"branch_id": b.branch_id, "branch_name": b.branch_name} for b in data.branches]
    products: dict[int, dict] = {}
    for branch in data.branches:
        for item in branch.items:
            row = products.setdefault(
                item.product_id,
                {
                    "product_code": item.product_code,
                    "category_name": item.category_name,
                    "description": item.product_description,
                    "cost_price": None,
                    "price_is_estimated": False,
                    "quantities": {},
                    "total": 0.0,
                },
            )
            row["quantities"][branch.branch_id] = item.quantity
            row["total"] += item.quantity
            if row["cost_price"] is None and item.effective_price is not None:
                row["cost_price"] = item.effective_price
                row["price_is_estimated"] = item.price_is_estimated

    rows = sorted(products.values(), key=lambda r: (r["category_name"], r["description"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Orders by Branch"

    header = ["Product Code", "Category", "Description", "Cost Price", "Price Basis"] + [
        b["branch_name"] for b in branches
    ] + ["Total"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    qty_fill = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")
    estimated_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    branch_totals = {b["branch_id"]: 0.0 for b in branches}
    grand_total = 0.0

    for row in rows:
        price_basis = "Estimated (older date)" if row["price_is_estimated"] else ("Confirmed" if row["cost_price"] is not None else "")
        values = [row["product_code"], row["category_name"], row["description"], row["cost_price"], price_basis]
        for b in branches:
            values.append(row["quantities"].get(b["branch_id"]))
        values.append(row["total"])
        ws.append(values)
        excel_row = ws.max_row
        if row["price_is_estimated"]:
            ws.cell(row=excel_row, column=4).fill = estimated_fill
            ws.cell(row=excel_row, column=5).fill = estimated_fill
        for i, b in enumerate(branches):
            qty = row["quantities"].get(b["branch_id"])
            if qty is not None:
                ws.cell(row=excel_row, column=6 + i).fill = qty_fill
                branch_totals[b["branch_id"]] += qty
        grand_total += row["total"]

    totals_row = ["", "", "Total", "", ""] + [branch_totals[b["branch_id"]] for b in branches] + [grand_total]
    ws.append(totals_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    widths = [12, 14, 30, 10, 18] + [12] * len(branches) + [10]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "F2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    date_label = data.delivery_date.isoformat() if data.delivery_date else "no-orders"
    filename = f"orders-by-branch-{date_label}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/admin", response_model=SupplierOrderAdminOut)
def admin_get_supplier_order(
    supplier_id: int,
    delivery_date: date,
    admin: User = Depends(require_roles("ADMIN")),
    db: Session = Depends(get_db),
):
    """Current order lines for one supplier on one delivery date — used to prefill the builder when editing."""
    return supplier_order_service.get_supplier_order(db, supplier_id, delivery_date)


@router.get("/admin/assigned", response_model=dict[int, float])
def admin_assigned_quantities(
    delivery_date: date,
    exclude_supplier_id: int | None = None,
    admin: User = Depends(require_roles("ADMIN")),
    db: Session = Depends(get_db),
):
    """Total quantity already given to suppliers per product, for one delivery date — pass the supplier
    currently being edited as exclude_supplier_id so their own lines don't count against themselves."""
    return supplier_order_service.get_assigned_quantities(db, delivery_date, exclude_supplier_id)


@router.get("/admin/prices", response_model=list[SupplierPricePreviewOut])
def admin_supplier_price_preview(
    supplier_id: int,
    delivery_date: date,
    admin: User = Depends(require_roles("ADMIN")),
    db: Session = Depends(get_db),
):
    """This supplier's resolved price for every product they've ever quoted — shown as a reference
    column on the Order Builder while Admin is still deciding quantities, before saving any lines."""
    resolved = supplier_order_service.get_supplier_price_preview(db, supplier_id, delivery_date)
    return [
        SupplierPricePreviewOut(product_id=pid, price=r.price, is_estimated=r.is_estimated, as_of=r.as_of)
        for pid, r in resolved.items()
    ]


@router.put("/admin", response_model=SupplierOrderAdminOut)
def admin_set_supplier_order(
    supplier_id: int,
    delivery_date: date,
    payload: SetSupplierOrderRequest,
    admin: User = Depends(require_roles("ADMIN")),
    db: Session = Depends(get_db),
):
    """
    Replaces the full set of order lines for this supplier on this
    delivery date (send the complete list each time, not a diff) — e.g.
    Malabe: Avocado 200, Pineapple 100; Kalubovila: Avocado 300, Pineapple 200.
    """
    return supplier_order_service.set_supplier_order(db, admin, supplier_id, delivery_date, payload)
