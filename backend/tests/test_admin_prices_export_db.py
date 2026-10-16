"""
DB-backed tests for Admin's Supplier Prices Excel download —
pricing_service.build_admin_prices_excel.
"""
import io
from datetime import datetime, timedelta

from openpyxl import load_workbook

from app.models.market_reference_price import MarketReferencePrice
from app.services import pricing_service
from tests.test_price_approval_db import _adjust_all, ctx  # noqa: F401


def _sheet(content: bytes):
    return load_workbook(io.BytesIO(content)).active


def _rows_by_code(ws):
    return {ws.cell(row=r, column=1).value: [c.value for c in ws[r]] for r in range(7, ws.max_row + 1)}


def test_excel_has_keells_mkt_and_each_suppliers_prices(db_session, ctx):
    today = datetime.now(pricing_service.BUSINESS_TZ).date()
    first, second = ctx["rows"]
    # Today's Keells price for the first item; only an older one for the second.
    db_session.add(MarketReferencePrice(product_id=first.product_id, source="KEELLS", price=200, delivery_date=today))
    old = today - timedelta(days=3)
    db_session.add(MarketReferencePrice(product_id=second.product_id, source="KEELLS", price=150, delivery_date=old))
    db_session.commit()
    _adjust_all(db_session, ctx)  # 100 -> 90, 110 -> 95

    content, used = pricing_service.build_admin_prices_excel(db_session, ctx["delivery"])
    assert used == ctx["delivery"]
    ws = _sheet(content)
    assert ws.cell(row=5, column=7).value == ctx["supplier"].supplier_name
    assert [ws.cell(row=6, column=c).value for c in (7, 8, 9)] == ["Supplier Price", "Adjusted Price", "Status"]

    rows = _rows_by_code(ws)
    codes = {r.product_id: db_session.get(pricing_service.Product, r.product_id).product_code for r in ctx["rows"]}
    a, b = rows[codes[first.product_id]], rows[codes[second.product_id]]
    assert a[3:6] == [200, datetime.combine(today, datetime.min.time()), 140]
    assert a[6:9] == [100, 90, "Draft"]
    # Carried forward from the last Keells entry, with that entry's date.
    assert b[3:6] == [150, datetime.combine(old, datetime.min.time()), 105]
    assert b[6:9] == [110, 95, "Draft"]


def test_excel_without_keells_or_prices(db_session, ctx, make_supplier):
    content, _ = pricing_service.build_admin_prices_excel(db_session, ctx["delivery"])
    rows = _rows_by_code(_sheet(content))
    assert all(r[3] is None and r[5] is None for r in rows.values())

    # A supplier with no quotes: header only, no product rows.
    content, _ = pricing_service.build_admin_prices_excel(db_session, ctx["delivery"], make_supplier().id)
    assert _sheet(content).max_row == 6


def test_export_endpoint_is_admin_only(client, ctx, auth_headers):
    url = f"/api/v1/pricing/admin/export?delivery_date={ctx['delivery'].isoformat()}"
    assert client.get(url, headers=auth_headers(ctx["supplier_user"])).status_code == 403
    res = client.get(url, headers=auth_headers(ctx["admin"]))
    assert res.status_code == 200
    assert res.content[:2] == b"PK"
    assert f"supplier-prices-{ctx['delivery'].isoformat()}.xlsx" in res.headers["content-disposition"]
