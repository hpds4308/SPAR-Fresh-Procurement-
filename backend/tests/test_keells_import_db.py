"""
DB-backed tests for the Keells price Excel import (keells_import_service)
and for the Keells sync never touching saved prices when it fails.
"""
import io
from datetime import date, datetime, timedelta

import pytest
from openpyxl import Workbook, load_workbook

from app.core.errors import ValidationFailedError
from app.models.market_reference_price import MarketReferencePrice
from app.models.product import Product
from app.services import keells_import_service as ki
from app.services import keells_scrape_service as ks
from app.services import pricing_service

HEADERS = ["DC Code", "System Name", "Keells Product Name", "Category", "Price", "Unit", "Numeric Price", "Scraped At"]
DAY = date(2026, 10, 16)
SCRAPED = datetime(2026, 10, 16, 7, 30)


def _row(code, name, keells, price, unit="KG", numeric="same", category="Local Fruits"):
    text = f"Rs {price:,.2f} /  {unit}" if isinstance(price, (int, float)) else price
    return [code, name, keells, category, text, unit, price if numeric == "same" else numeric, SCRAPED]


GOOD_ROWS = [
    _row("4503253", "AVOCADO", "Avocado", 290),
    _row("4503272", "MANGO TJC GRADE 1", "Mango - Tjc", 1650),
    _row("4503775", "COCONUTS", "Coconut", 166, unit="Unit", category="Fresh Vegetables"),
    _row("4503285", "PAPAYA", "Papaya", 270),
]


def _xlsx(rows, headers=HEADERS) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Required Products"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def keells_products(db_session, make_product):
    """The products the file rows point at (get-or-create: the test DB may already hold them)."""
    made = {}
    for code, name in [
        ("4503253", "AVOCADO"),
        ("4503272", "MANGO TJC GRADE 1"),
        ("4503775", "COCONUTS"),
        ("4503285", "PAPAYA"),
        ("4503286", "PAPAYA TANNING"),
    ]:
        product = db_session.query(Product).filter(Product.product_code == code).first()
        if product:
            product.description, product.status = name, "ACTIVE"
            db_session.commit()
        else:
            product = make_product(product_code=code, description=name)
        made[code] = product
    return made


def _prices(db_session, day=DAY):
    rows = (
        db_session.query(MarketReferencePrice)
        .filter(MarketReferencePrice.source == "KEELLS", MarketReferencePrice.delivery_date == day)
        .all()
    )
    return {r.product_id: float(r.price) for r in rows}


def _errors(preview):
    return {l["row_number"]: l["error"] for l in preview["lines"] if l["error"]}


# ---- Valid imports ----


def test_valid_file_imports_every_row(db_session, keells_products, make_user):
    admin = make_user(role="ADMIN")
    preview = ki.preview_keells_excel(db_session, DAY, _xlsx(GOOD_ROWS))
    assert (preview["valid_count"], preview["error_count"], preview["existing_count"]) == (4, 0, 0)
    assert _prices(db_session) == {}  # preview saves nothing

    result = ki.import_keells_excel(db_session, admin, DAY, _xlsx(GOOD_ROWS), "keells.xlsx")
    assert (result["imported"], result["replaced"]) == (4, 0)
    p = keells_products
    assert _prices(db_session) == {
        p["4503253"].id: 290.0,
        p["4503272"].id: 1650.0,
        p["4503775"].id: 166.0,
        p["4503285"].id: 270.0,
    }


def test_reuploading_replaces_without_duplicates(db_session, keells_products, make_user):
    admin = make_user(role="ADMIN")
    ki.import_keells_excel(db_session, admin, DAY, _xlsx(GOOD_ROWS))
    result = ki.import_keells_excel(db_session, admin, DAY, _xlsx(GOOD_ROWS))
    assert (result["imported"], result["replaced"]) == (4, 4)
    count = (
        db_session.query(MarketReferencePrice)
        .filter(MarketReferencePrice.source == "KEELLS", MarketReferencePrice.delivery_date == DAY)
        .count()
    )
    assert count == 4

    # A smaller file replaces the day: dropped products lose that day's price.
    ki.import_keells_excel(db_session, admin, DAY, _xlsx(GOOD_ROWS[:1]))
    assert _prices(db_session) == {keells_products["4503253"].id: 290.0}


def test_import_leaves_other_dates_and_sources_alone(db_session, keells_products, make_user):
    admin = make_user(role="ADMIN")
    avocado = keells_products["4503253"].id
    other_day = DAY - timedelta(days=2)
    pricing_service.set_reference_price(db_session, admin, avocado, other_day, 250, "KEELLS")
    pricing_service.set_reference_price(db_session, admin, avocado, DAY, 120, "LOCAL_MARKET")

    ki.import_keells_excel(db_session, admin, DAY, _xlsx(GOOD_ROWS))
    assert _prices(db_session, other_day) == {avocado: 250.0}
    local = db_session.query(MarketReferencePrice).filter_by(source="LOCAL_MARKET", delivery_date=DAY).one()
    assert float(local.price) == 120.0


def test_cell_types_codes_as_numbers_and_comma_prices(db_session, keells_products):
    rows = [
        [4503253, "Avocado", "Avocado", "Local Fruits", "Rs 1,290.00 / KG", "kg", "1290", SCRAPED],
        [4503272.0, "MANGO TJC GRADE 1", "Mango - Tjc", "Local Fruits", None, "KG", 1650.5, SCRAPED],
        ["4503285", "PAPAYA", "Papaya", "Local Fruits", "Rs 270.00 / KG", "KG", None, SCRAPED],
    ]
    preview = ki.preview_keells_excel(db_session, DAY, _xlsx(rows))
    assert preview["error_count"] == 0, _errors(preview)
    assert [l["price"] for l in preview["lines"]] == [1290.0, 1650.5, 270.0]


def test_not_in_file_lists_missing_keells_products(db_session, keells_products):
    preview = ki.preview_keells_excel(db_session, DAY, _xlsx(GOOD_ROWS[:2]))
    missing = {m["dc_code"] for m in preview["not_in_file"]}
    assert {"4503775", "4503285"} <= missing
    assert "4503286" not in missing  # PAPAYA TANNING isn't compared with Keells at all


# ---- Rejected files and rows ----


def test_not_an_excel_file(db_session):
    with pytest.raises(ValidationFailedError):
        ki.preview_keells_excel(db_session, DAY, b"DC Code,Price\n4503253,290\n")
    with pytest.raises(ValidationFailedError):
        ki.preview_keells_excel(db_session, DAY, b"PK\x03\x04 broken zip")
    with pytest.raises(ValidationFailedError) as e:
        ki.preview_keells_excel(db_session, DAY, bytes.fromhex("d0cf11e0a1b11ae1") + b"\0" * 64)
    assert ".xls" in e.value.message


def test_missing_headers(db_session):
    headers = ["DC Code", "System Name", "Keells Product Name", "Category", "Price", "Scraped At"]
    with pytest.raises(ValidationFailedError) as e:
        ki.preview_keells_excel(db_session, DAY, _xlsx([], headers=headers))
    assert "Unit" in e.value.message and "Numeric Price" in e.value.message
    with pytest.raises(ValidationFailedError):
        ki.preview_keells_excel(db_session, DAY, _xlsx([], headers=["Product", "Cost"]))


def test_bad_rows_are_reported_row_by_row(db_session, keells_products):
    rows = [
        _row("4503253", "AVOCADO", "Avocado", 290),  # row 2: fine
        _row("9999999", "DRAGONFRUIT X", "Dragon", 900),  # 3: not on the Keells list
        _row("4503253", "AVOCADO", "Avocado", 300),  # 4: duplicate
        _row("4503272", "MANGO TJC GRADE 1", "Mango - Tjc", "abc", numeric=None),  # 5: unreadable price
        _row("4503775", "COCONUTS", "Coconut", None, unit="Unit", numeric=None),  # 6: missing price
        _row("4503285", "PAPAYA", "Papaya", 270, unit="Bunch"),  # 7: unknown unit
    ]
    errors = _errors(ki.preview_keells_excel(db_session, DAY, _xlsx(rows)))
    assert set(errors) == {3, 4, 5, 6, 7}
    assert "isn't on the list" in errors[3]
    assert "more than once" in errors[4] and "row 2" in errors[4]
    assert "isn't a Keells price" in errors[5]
    assert errors[6] == "Price is missing."
    assert "KG or Unit" in errors[7]


@pytest.mark.parametrize(
    "row, expected",
    [
        (_row("4503253", "AVOCADO", "Avocado", "Rs 290.00 /  KG", numeric=320), "don't agree"),
        (_row("4503253", "AVOCADO", "Avocado", "Rs 290.00 /  Unit", numeric=290), "per Unit"),
        (_row("4503253", "AVOCADO", "Avocado", 0), "greater than zero"),
        (_row("4503253", "AVOCADO", "Avocado", None, numeric=-5), "greater than zero"),
        (_row("4503253", "AVOCADO", "Avocado", "Rs -5.00 / KG", numeric=None), "isn't a Keells price"),
        (_row("4503253", "AVOCADO", "Avocado", "Rs 290.00 / KG", numeric="#N/A"), "Excel error"),
        (_row("4503253", "AVOCADO", "Avocado", 290, unit=None), "Unit is missing"),
        (_row(None, "AVOCADO", "Avocado", 290), "DC Code is missing"),
        (_row("4503253", None, "Avocado", 290), "System Name is missing"),
    ],
)
def test_invalid_prices_units_and_values(db_session, keells_products, row, expected):
    errors = _errors(ki.preview_keells_excel(db_session, DAY, _xlsx([row])))
    assert expected in errors[2]


def test_papaya_tanning_never_maps_to_papaya(db_session, keells_products):
    rows = [
        _row("4503286", "PAPAYA TANNING", "Papaya", 270),  # not compared with Keells
        _row("4503285", "PAPAYA TANNING", "Papaya", 270),  # PAPAYA's code with the wrong name
    ]
    errors = _errors(ki.preview_keells_excel(db_session, DAY, _xlsx(rows)))
    assert "isn't on the list" in errors[2]
    assert "doesn't match" in errors[3]


def test_any_bad_row_blocks_the_whole_import(db_session, keells_products, make_user):
    admin = make_user(role="ADMIN")
    avocado = keells_products["4503253"].id
    pricing_service.set_reference_price(db_session, admin, avocado, DAY, 280, "KEELLS")

    bad = GOOD_ROWS + [_row("4503253", "AVOCADO", "Avocado", 300)]  # duplicate
    with pytest.raises(ValidationFailedError) as e:
        ki.import_keells_excel(db_session, admin, DAY, _xlsx(bad))
    assert "nothing was imported" in e.value.message
    assert _prices(db_session) == {avocado: 280.0}


def test_database_failure_rolls_back(db_session, keells_products, make_user, monkeypatch):
    admin = make_user(role="ADMIN")
    avocado = keells_products["4503253"].id
    pricing_service.set_reference_price(db_session, admin, avocado, DAY, 280, "KEELLS")

    def broken_commit():
        raise RuntimeError("database went away")

    monkeypatch.setattr(db_session, "commit", broken_commit)
    with pytest.raises(RuntimeError):
        ki.import_keells_excel(db_session, admin, DAY, _xlsx(GOOD_ROWS))
    monkeypatch.undo()
    db_session.expire_all()
    assert _prices(db_session) == {avocado: 280.0}


# ---- Failed sync leaves prices alone ----


class _FakeBrowser:
    def new_page(self):
        return object()

    def close(self):
        pass


class _FakePlaywright:
    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    @property
    def chromium(self):
        return self

    def launch(self, **kwargs):
        return _FakeBrowser()


def test_blocked_sync_keeps_existing_prices(db_session, keells_products, make_user, monkeypatch):
    import playwright.sync_api

    admin = make_user(role="ADMIN")
    avocado = keells_products["4503253"].id
    pricing_service.set_reference_price(db_session, admin, avocado, DAY, 280, "KEELLS")

    def blocked(*args, **kwargs):
        raise ks.KeellsScrapeError(ks.BLOCKED_MESSAGE)

    monkeypatch.setattr(playwright.sync_api, "sync_playwright", lambda: _FakePlaywright())
    monkeypatch.setattr(ks, "_scrape_category", blocked)
    with pytest.raises(ks.KeellsScrapeError):
        ks.run_scrape(db_session, admin, DAY)
    assert _prices(db_session) == {avocado: 280.0}


def test_sync_endpoint_returns_502_with_cors_headers(client, db_session, make_user, auth_headers, monkeypatch):
    from app.core.config import settings

    def blocked(*args, **kwargs):
        raise ks.KeellsScrapeError(ks.BLOCKED_MESSAGE)

    monkeypatch.setattr(ks, "run_scrape", blocked)
    origin = settings.CORS_ORIGINS[0]
    res = client.post(
        f"/api/v1/pricing/reference/keells-sync?delivery_date={DAY.isoformat()}",
        headers={**auth_headers(make_user(role="ADMIN")), "Origin": origin},
    )
    assert res.status_code == 502
    assert res.json() == {"detail": ks.BLOCKED_MESSAGE}
    assert res.headers.get("access-control-allow-origin") == origin


# ---- Endpoints, permissions, and where the prices show up ----


def test_endpoints_admin_only_and_prices_flow_to_supplier_prices(
    client, db_session, keells_products, make_user, auth_headers
):
    today = datetime.now(pricing_service.BUSINESS_TZ).date()
    files = {"file": ("keells-required-products.xlsx", _xlsx(GOOD_ROWS), "application/octet-stream")}
    for path in ("preview", "import"):
        url = f"/api/v1/pricing/reference/keells-excel/{path}?delivery_date={today.isoformat()}"
        for role in ("BRANCH", "SUPPLIER"):
            assert client.post(url, files=files, headers=auth_headers(make_user(role=role))).status_code == 403
        assert client.post(url, files=files).status_code == 401

    admin_h = auth_headers(make_user(role="ADMIN"))
    preview = client.post(
        f"/api/v1/pricing/reference/keells-excel/preview?delivery_date={today.isoformat()}", files=files, headers=admin_h
    )
    assert preview.status_code == 200 and preview.json()["valid_count"] == 4
    res = client.post(
        f"/api/v1/pricing/reference/keells-excel/import?delivery_date={today.isoformat()}", files=files, headers=admin_h
    )
    assert res.status_code == 200 and res.json()["imported"] == 4

    # Keells Prices page.
    listed = client.get(f"/api/v1/pricing/reference?delivery_date={today.isoformat()}", headers=admin_h).json()
    assert {r["product_id"]: r["price"] for r in listed}[keells_products["4503272"].id] == 1650.0

    # Supplier Prices reads today's Keells price (else the latest one); its
    # Excel shows MKT Price = Keells x 70%.
    latest = client.get("/api/v1/pricing/reference/last", headers=admin_h).json()
    assert {r["product_id"]: r["price"] for r in latest}[keells_products["4503253"].id] == 290.0
    assert pricing_service.MKT_FACTOR == 0.7


def test_supplier_prices_excel_uses_imported_keells_price(db_session, keells_products, make_user, make_supplier):
    from app.models.pricing import SupplierPrice

    today = datetime.now(pricing_service.BUSINESS_TZ).date()
    admin = make_user(role="ADMIN")
    supplier = make_supplier()
    supplier_user = make_user(role="SUPPLIER", supplier=supplier)
    mango = keells_products["4503272"]
    delivery = today + timedelta(days=3)
    db_session.add(
        SupplierPrice(
            supplier_id=supplier.id, product_id=mango.id, delivery_date=delivery, price=1200, unit_code="KG",
            submitted_by=supplier_user.id,
        )
    )
    db_session.commit()
    ki.import_keells_excel(db_session, admin, today, _xlsx(GOOD_ROWS))

    content, _ = pricing_service.build_admin_prices_excel(db_session, delivery)
    ws = load_workbook(io.BytesIO(content)).active
    row = next(r for r in ws.iter_rows(min_row=7, values_only=True) if str(r[0]) == mango.product_code)
    assert row[3] == 1650 and row[5] == 1155  # Keells, MKT = 1650 x 0.7
