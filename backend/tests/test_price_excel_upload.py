"""
DB-backed tests for the supplier price Excel template download and
upload/preview endpoints.
"""
import io

import openpyxl
import pytest

from app.models.system import SystemSetting

HEADERS = ("Product Code", "Product", "Price")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx_bytes(rows, headers=HEADERS, title_rows=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    for t in title_rows:
        ws.append(list(t))
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def supplier_user(make_supplier, make_user):
    return make_user(role="SUPPLIER", supplier=make_supplier())


@pytest.fixture()
def open_price_window(db_session, make_user):
    admin = make_user(role="ADMIN")
    db_session.add(SystemSetting(key="supplier_price_deadline", value="23:59", updated_by=admin.id))
    db_session.add(SystemSetting(key="supplier_price_days", value="0,1,2,3,4,5,6", updated_by=admin.id))
    db_session.commit()


def _upload(client, headers, content, filename="prices.xlsx"):
    return client.post("/api/v1/pricing/preview-excel", headers=headers, files={"file": (filename, content, XLSX)})


def test_template_lists_every_active_product(client, auth_headers, supplier_user, make_product):
    a = make_product(description="Apple")
    b = make_product(description="Banana")
    inactive = make_product(description="Gone", status="INACTIVE")

    res = client.get("/api/v1/pricing/excel-template", headers=auth_headers(supplier_user))
    assert res.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
    assert len(ws._images) == 1  # SPAR logo
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    assert rows[0] == ("Product Code", "Category", "Product", "Price")
    codes = {r[0] for r in rows[1:]}
    assert {a.product_code, b.product_code} <= codes
    assert inactive.product_code not in codes
    assert all(r[1] for r in rows[1:])  # category filled in
    assert all(r[3] is None for r in rows[1:])


def test_branch_cannot_download_template(client, auth_headers, make_user, make_branch):
    branch_user = make_user(role="BRANCH", branch=make_branch())
    res = client.get("/api/v1/pricing/excel-template", headers=auth_headers(branch_user))
    assert res.status_code == 403


def test_filled_template_round_trips(client, auth_headers, supplier_user, make_product):
    a = make_product(description="Apple")
    b = make_product(description="Banana")

    template = client.get("/api/v1/pricing/excel-template", headers=auth_headers(supplier_user)).content
    wb = openpyxl.load_workbook(io.BytesIO(template))
    ws = wb.active
    for row in ws.iter_rows(min_row=5):
        if row[0].value == a.product_code:
            row[3].value = 250.5
    buf = io.BytesIO()
    wb.save(buf)

    res = _upload(client, auth_headers(supplier_user), buf.getvalue())
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["valid_line_count"] == 1
    assert body["error_count"] == 0
    assert body["blank_price_count"] >= 1  # Banana left blank
    line = body["lines"][0]
    assert line["product_id"] == a.id
    assert line["price"] == 250.5
    assert b.id not in [l["product_id"] for l in body["lines"]]


def test_price_text_and_errors(client, auth_headers, supplier_user, make_product):
    a = make_product()
    b = make_product()
    c = make_product()
    content = _xlsx_bytes(
        [
            (a.product_code, "x", "Rs. 1,250.00"),
            (b.product_code, "x", "abc"),
            (c.product_code, "x", 0),
            ("NOPE", "x", 100),
            (a.product_code, "x", 99),
        ]
    )
    body = _upload(client, auth_headers(supplier_user), content).json()
    by_row = {l["row_number"]: l for l in body["lines"]}
    assert by_row[2]["price"] == 1250.0 and by_row[2]["error"] is None
    assert "not a valid price" in by_row[3]["error"]
    assert "greater than zero" in by_row[4]["error"]
    assert "Unknown" in by_row[5]["error"]
    assert "Duplicate" in by_row[6]["error"]
    assert body["valid_line_count"] == 1
    assert body["error_count"] == 4


def test_header_below_title_rows_and_numeric_codes(client, auth_headers, supplier_user, make_product):
    p = make_product(product_code="1001")
    content = _xlsx_bytes([(1001, "Thing", 75)], title_rows=[("ABC Suppliers price list",), ()])
    body = _upload(client, auth_headers(supplier_user), content).json()
    assert body["valid_line_count"] == 1
    assert body["lines"][0]["product_id"] == p.id


def test_missing_columns_rejected(client, auth_headers, supplier_user):
    res = _upload(client, auth_headers(supplier_user), _xlsx_bytes([("A", "B")], headers=("Product Code", "Product")))
    assert res.status_code == 422
    assert "header row" in res.json()["detail"]


def test_garbage_file_rejected(client, auth_headers, supplier_user):
    res = _upload(client, auth_headers(supplier_user), b"not a spreadsheet")
    assert res.status_code == 422


def test_branch_cannot_upload(client, auth_headers, make_user, make_branch):
    branch_user = make_user(role="BRANCH", branch=make_branch())
    res = _upload(client, auth_headers(branch_user), _xlsx_bytes([]))
    assert res.status_code == 403


def test_preview_then_submit_saves_prices(client, auth_headers, supplier_user, make_product, open_price_window):
    a = make_product()
    headers = auth_headers(supplier_user)
    body = _upload(client, headers, _xlsx_bytes([(a.product_code, "x", 180)])).json()
    entries = [{"product_id": l["product_id"], "price": l["price"]} for l in body["lines"] if not l["error"]]
    res = client.post("/api/v1/pricing", headers=headers, json={"prices": entries})
    assert res.status_code == 200, res.text
    assert res.json()[0]["price"] == 180.0
